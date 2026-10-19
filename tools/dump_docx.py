"""导出 docx 的结构：段落、样式、表格，用于对照模板格式。
用法：python dump_docx.py <文件路径>
"""
import sys
from docx import Document
from docx.shared import Pt


def fmt(p):
    """简述段落格式"""
    pf = p.paragraph_format
    bits = []
    if pf.alignment is not None:
        bits.append(f"align={pf.alignment}")
    if pf.first_line_indent is not None:
        bits.append(f"indent={pf.first_line_indent.pt:.0f}pt")
    if pf.line_spacing is not None:
        bits.append(f"line={pf.line_spacing}")
    sizes = {r.font.size.pt for r in p.runs if r.font.size}
    if sizes:
        bits.append("size=" + "/".join(f"{s:g}" for s in sorted(sizes)))
    fonts = {r.font.name for r in p.runs if r.font.name}
    if fonts:
        bits.append("font=" + ",".join(sorted(fonts)))
    if any(r.bold for r in p.runs):
        bits.append("bold")
    return " ".join(bits)


def main(path):
    doc = Document(path)
    print(f"=== 段落总数 {len(doc.paragraphs)}，表格 {len(doc.tables)} ===\n")
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if not text:
            continue
        print(f"[{i:03d}] style={p.style.name!r} | {fmt(p)}")
        print(f"      {text}")

    for ti, tbl in enumerate(doc.tables):
        print(f"\n=== 表格 {ti}: {len(tbl.rows)} 行 x {len(tbl.columns)} 列 ===")
        for ri, row in enumerate(tbl.rows):
            cells = []
            for c in row.cells:
                cells.append(c.text.replace("\n", "⏎").strip())
            print(f"  r{ri}: " + " | ".join(cells))

    # 节面设置
    for si, sec in enumerate(doc.sections):
        print(f"\n=== 节 {si}: 页面 {sec.page_width.cm:.1f}x{sec.page_height.cm:.1f}cm "
              f"边距 上{sec.top_margin.cm:.1f} 下{sec.bottom_margin.cm:.1f} "
              f"左{sec.left_margin.cm:.1f} 右{sec.right_margin.cm:.1f} ===")


if __name__ == "__main__":
    main(sys.argv[1])
